"""Build DNNP paste CSVs in the PostAward SharePoint column order."""

import csv
import re
import uuid
from datetime import date, datetime
from pathlib import Path

import openpyxl

XLSX = Path(
    r"c:\Users\gabri\Downloads\DNNP_DA1_MASTER CO-CLAIMS REGISTER_Updated Sep 4 2025 (3).xlsx"
)
OUT_DIR = Path(__file__).resolve().parents[1] / "seed_data" / "DNNP"
PROJECT = "A02180"
NAMESPACE = uuid.UUID("6f1b0a3e-9c2d-4e5f-a1b2-0d4e6f8a9c01")

SUB_COLUMNS = [
    "Title",
    "subcontractnumber",
    "packageidentifier",
    "projectidentifier",
    "ponumber",
    "vendor",
    "scope",
    "bundle",
    "contracttype",
    "originalvalue",
    "startdate",
    "completiondate",
    "responsibleca",
    "backupca",
    "currentlyactive",
    "holdback",
    "closeoutstatus",
    "invoicedtodate",
    "status",
    "sapvendornumber",
    "revisedenddate",
    "currency",
    "indigenouscontribution",
    "worktype",
    "operationalowner",
    "operationalowneremail",
]

CHANGE_COLUMNS = [
    "Title",
    "subcontractid",
    "projectidentifier",
    "refnumber",
    "changedate",
    "description",
    "changetype",
    "reason",
    "value",
    "status",
    "aeconref",
    "datereceived",
    "datefullyexecuted",
    "initialvalue",
    "finalvalue",
    "budget",
]

SKIP_SHEETS = {
    "Dropdowns",
    "Detail1",
    "Pivot_Summary",
    "Sheet2",
    "Sub Summary",
    "CO Pick Number",
    "CD Pick Number",
}

CHANGE_HEADER = "SUBCONTRACTOR CHANGE ORDER REQUESTS NO.:"


def clean(value) -> str:
    if value is None:
        return ""
    text = str(value).replace("\u202f", " ").replace("\xa0", " ").strip()
    return re.sub(r"\s+", " ", text)


def compact_id(value) -> str:
    return re.sub(r"\s+", "", clean(value))


def as_number(value):
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = clean(value).replace(",", "").replace("$", "")
    if text in {"", "-", "—"}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def money(value) -> str:
    number = as_number(value)
    if number is None:
        return ""
    return f"{number:.2f}"


def ref_code(value) -> str:
    if isinstance(value, str):
        return clean(value)
    if isinstance(value, bool) or value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        number = int(value)
        return f"{number:03d}" if number < 1000 else str(number)
    if isinstance(value, int):
        return f"{value:03d}" if value < 1000 else str(value)
    return clean(value)


def as_text_id(value) -> str:
    number = as_number(value)
    if number is None:
        return clean(value)
    if number.is_integer():
        return str(int(number))
    return clean(value)


def as_date(value) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.strftime("%m/%d/%Y")
    if isinstance(value, date):
        return value.strftime("%m/%d/%Y")
    text = clean(value)
    if text.lower() in {"", "tbd", "n/a", "na", "-", "—"}:
        return ""
    for pattern in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%d-%b-%y", "%d-%b-%Y", "%d-%B-%Y"):
        try:
            return datetime.strptime(text, pattern).strftime("%m/%d/%Y")
        except ValueError:
            continue
    return ""


def yes_no(value) -> str:
    text = clean(value).lower()
    if text in {"yes", "y", "true", "1"}:
        return "Yes"
    if text in {"no", "n", "false", "0"}:
        return "No"
    return ""


def holdback_and_closeout(raw: str) -> tuple[str, str]:
    text = clean(raw).lower()
    if not text:
        return "Applicable", "In progress"
    if "not applicable" in text:
        return "Not applicable", "Close-out pending"
    if "full hb processed" in text or "close-out done" in text:
        return "Applicable", "Closed"
    if "hb" in text and "to be processed" in text:
        return "Applicable", "Holdback to process"
    if "to be close" in text:
        holdback = "Not applicable" if "no hb" in text else "Applicable"
        return holdback, "Close-out pending"
    return "", ""


def work_type(scope: str, bundle: str) -> str:
    scope_text = scope.lower()
    if "consult" in scope_text:
        return "Consulting"
    if "supply only" in scope_text or scope_text.startswith("supply"):
        return "Supply only"
    if clean(bundle).upper() == "BOP":
        return "Boots on ground"
    return ""


def change_status(raw: str) -> str:
    text = clean(raw).lower()
    if text == "approved":
        return "Approved"
    if text.startswith("pending"):
        return "Pending"
    if text.startswith("potential"):
        return "Potential"
    if text.startswith("rejected"):
        return "Rejected"
    if "cancel" in text:
        return "Canceled"
    return ""


def change_type(raw: str) -> str:
    text = clean(raw).lower()
    if "directive" in text:
        return "Change Directive"
    if "claim" in text:
        return "Claim"
    if "change order" in text or text == "co":
        return "Change Order"
    return ""


def reason_from(row: list) -> str:
    """Reason columns sit at indexes 8-11: design, scope gap, variation, acceleration."""
    labeled = [
        ("Design Change", 8),
        ("Scope Gap", 9),
        ("Variation", 10),
        ("Acceleration", 11),
    ]
    scored = []
    for label, index in labeled:
        if index >= len(row):
            continue
        amount = as_number(row[index])
        if amount not in (None, 0):
            scored.append((abs(amount), label))
    if not scored:
        return ""
    scored.sort(reverse=True)
    return scored[0][1]


def display_title(subcontract_no: str, number: str) -> str:
    subcontract_no = compact_id(subcontract_no)
    number = clean(number)
    if subcontract_no and number:
        tail = subcontract_no.rsplit("-", 1)[-1]
        same_tail = tail.lower() == number.lower() or tail.lstrip("0") == number.lstrip("0")
        if same_tail:
            return subcontract_no
        return f"{subcontract_no}-{number}"
    return subcontract_no or number


def package_id(title: str, po_number: str) -> str:
    key = compact_id(title) or as_text_id(po_number)
    if po_number and po_number not in {"0"}:
        key = f"{key}|{as_text_id(po_number)}"
    return str(uuid.uuid5(NAMESPACE, "DNNP|" + key))


def find_header(rows: list[list], label: str):
    for index, row in enumerate(rows):
        for cell in row[:8]:
            if clean(cell) == label:
                return index
    return None


def cell_after_label(rows: list[list], label: str) -> str:
    for row in rows[:20]:
        for index, cell in enumerate(row[:6]):
            if clean(cell).rstrip(":").upper() == label.upper():
                for follower in row[index + 1 : index + 4]:
                    if clean(follower):
                        return clean(follower)
    return ""


def summary_rows(ws) -> list[dict]:
    rows = [list(row) for row in ws.iter_rows(max_col=40, values_only=True)]
    header_at = find_header(rows, "PO#")
    if header_at is None:
        raise SystemExit("Sub Summary is missing the PO# header")
    header = [clean(cell) for cell in rows[header_at]]
    index = {name: position for position, name in enumerate(header) if name}

    def take(row, name):
        position = index.get(name)
        if position is None or position >= len(row):
            return None
        return row[position]

    records = []
    for row in rows[header_at + 1 :]:
        vendor = clean(take(row, "VENDOR NAME"))
        po_number = as_text_id(take(row, "PO#"))
        subcontract_no = compact_id(take(row, "SUBCONTRACT#"))
        if vendor in {"", "0"} or vendor.upper() in {"VENDOR NAME", "TOTAL"}:
            continue
        if po_number in {"", "0"} and not subcontract_no:
            continue
        records.append(
            {
                "number": as_text_id(take(row, "NO.:")) or subcontract_no.rsplit("-", 1)[-1],
                "subcontract_no": subcontract_no,
                "po": po_number,
                "vendor": vendor,
                "scope": clean(take(row, "SCOPE")),
                "original": take(row, "ORGINAL CONTRACT"),
                "invoiced": take(row, "APPROVED INVOICE AMOUNT TO DATE (or Keyed in SAP)"),
                "ca": clean(take(row, "Primary Resp. CA")),
                "backup": clean(take(row, "Resp. CA (In case of Primary is Absence)")),
                "start": take(row, "Sub-contract Start Date"),
                "end": take(row, "Sub-contract Completion Date"),
                "active": take(row, "Sub-contract Currently Active (Yes/No)"),
                "bundle": clean(take(row, "Bundle")),
                "closeout": clean(take(row, "Sub-contract Close-out Status")),
            }
        )
    return records


def sheet_changes(ws) -> tuple[dict, list[dict]]:
    rows = [list(row) for row in ws.iter_rows(max_col=16, values_only=True)]
    header_at = None
    for index, row in enumerate(rows[:80]):
        if clean(row[0] if row else "") == CHANGE_HEADER:
            header_at = index
            break
    if header_at is None:
        return {}, []
    identity = {
        "vendor": cell_after_label(rows, "SUBCONTRACTOR"),
        "subcontract_no": compact_id(cell_after_label(rows, "SUBCONTRACT NO.")),
        "po": as_text_id(cell_after_label(rows, "PURCHASE ORDER NO.")),
        "scope": cell_after_label(rows, "SCOPE"),
        "original": cell_after_label(rows, "ORIGINAL CONTRACT VALUE"),
    }
    changes = []
    for row in rows[header_at + 1 :]:
        refnumber = ref_code(row[0] if row else "")
        description = clean(row[3] if len(row) > 3 else "")
        amount = as_number(row[2] if len(row) > 2 else None)
        if not refnumber and not description and amount is None:
            if changes:
                break
            continue
        if not refnumber or (amount is None and not description):
            continue
        status = change_status(row[5] if len(row) > 5 else "")
        kind = change_type(row[4] if len(row) > 4 else "")
        when = as_date(row[1] if len(row) > 1 else None)
        aecon = clean(row[12] if len(row) > 12 else "")
        if aecon in {"-", "—"}:
            aecon = ""
        changes.append(
            {
                "refnumber": refnumber,
                "date": when,
                "value": amount,
                "description": description,
                "changetype": kind,
                "status": status,
                "reason": reason_from(row),
                "aeconref": aecon,
                "raw_status": clean(row[5] if len(row) > 5 else ""),
                "raw_type": clean(row[4] if len(row) > 4 else ""),
            }
        )
    return identity, changes


def subcontract_record(item: dict) -> dict:
    holdback, closeout = holdback_and_closeout(item.get("closeout", ""))
    if not closeout and clean(item.get("closeout", "")):
        closeout = ""
        holdback = ""
    scope = item.get("scope", "")
    bundle = item.get("bundle", "")
    title = display_title(item.get("subcontract_no", ""), item.get("number", ""))
    if not title:
        title = item.get("vendor") or item.get("po") or ""
    return {
        "Title": title,
        "subcontractnumber": item.get("number") or title.rsplit("-", 1)[-1],
        "packageidentifier": package_id(title, item.get("po", "")),
        "projectidentifier": PROJECT,
        "ponumber": item.get("po", ""),
        "vendor": item.get("vendor", ""),
        "scope": scope,
        "bundle": bundle,
        "contracttype": "",
        "originalvalue": money(item.get("original")),
        "startdate": as_date(item.get("start")),
        "completiondate": as_date(item.get("end")),
        "responsibleca": item.get("ca", ""),
        "backupca": item.get("backup", ""),
        "currentlyactive": yes_no(item.get("active")),
        "holdback": holdback,
        "closeoutstatus": closeout,
        "invoicedtodate": money(item.get("invoiced")),
        "status": "",
        "sapvendornumber": "",
        "revisedenddate": "",
        "currency": "",
        "indigenouscontribution": "0",
        "worktype": work_type(scope, bundle),
        "operationalowner": "",
        "operationalowneremail": "",
        "_closeout_raw": item.get("closeout", ""),
    }


def write_csv(path: Path, columns: list[str], records: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore", lineterminator="\r\n")
        writer.writeheader()
        for record in records:
            writer.writerow({column: record.get(column, "") for column in columns})


def main() -> None:
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    summary = summary_rows(wb["Sub Summary"])
    by_sub = {}
    by_po = {}
    subcontracts = []
    for item in summary:
        record = subcontract_record(item)
        subcontracts.append(record)
        if item["subcontract_no"]:
            by_sub.setdefault(item["subcontract_no"], []).append(record)
        if item["po"] and item["po"] != "0":
            by_po[item["po"]] = record

    changes = []
    unmatched_sheets = []
    unknown_status = {}
    unknown_type = {}
    sheets_with_changes = 0
    for sheet_name in wb.sheetnames:
        if sheet_name in SKIP_SHEETS:
            continue
        identity, sheet_rows = sheet_changes(wb[sheet_name])
        if not sheet_rows and not identity:
            continue
        if sheet_rows:
            sheets_with_changes += 1
        po_key = identity.get("po", "")
        sub_key = identity.get("subcontract_no", "")
        record = by_po.get(po_key) if po_key and po_key != "0" else None
        if record is None:
            matches = by_sub.get(sub_key, [])
            if len(matches) == 1:
                record = matches[0]
        vendor_name = identity.get("vendor", "")
        po_name = identity.get("po", "")
        useful_identity = vendor_name not in {"", "0"} and (po_name not in {"", "0"} or identity.get("subcontract_no"))
        if record is None and useful_identity:
            created = subcontract_record(
                {
                    "number": (identity.get("subcontract_no") or "").rsplit("-", 1)[-1],
                    "subcontract_no": identity.get("subcontract_no", ""),
                    "po": identity.get("po", ""),
                    "vendor": identity.get("vendor", ""),
                    "scope": identity.get("scope", ""),
                    "original": identity.get("original"),
                    "invoiced": None,
                    "ca": "",
                    "backup": "",
                    "start": None,
                    "end": None,
                    "active": "",
                    "bundle": "",
                    "closeout": "",
                }
            )
            subcontracts.append(created)
            record = created
            if identity.get("subcontract_no"):
                by_sub.setdefault(identity["subcontract_no"], []).append(created)
            if identity.get("po") and identity["po"] != "0":
                by_po[identity["po"]] = created
            unmatched_sheets.append(sheet_name)
        if record is None:
            unmatched_sheets.append(sheet_name)
            continue
        seen = {}
        for change in sheet_rows:
            if change["raw_status"] and not change["status"]:
                unknown_status[change["raw_status"]] = unknown_status.get(change["raw_status"], 0) + 1
            if change["raw_type"] and not change["changetype"]:
                unknown_type[change["raw_type"]] = unknown_type.get(change["raw_type"], 0) + 1
            ref = change["refnumber"]
            seen[ref] = seen.get(ref, 0) + 1
            suffix = "" if seen[ref] == 1 else f"-{seen[ref]}"
            title = f"{ref}{suffix} {change['description']}".strip()
            value = money(change["value"])
            changes.append(
                {
                    "Title": title[:255],
                    "subcontractid": record["packageidentifier"],
                    "projectidentifier": PROJECT,
                    "refnumber": ref,
                    "changedate": change["date"],
                    "description": change["description"],
                    "changetype": change["changetype"],
                    "reason": change["reason"],
                    "value": value,
                    "status": change["status"],
                    "aeconref": change["aeconref"],
                    "datereceived": change["date"],
                    "datefullyexecuted": "",
                    "initialvalue": value,
                    "finalvalue": value if change["status"] == "Approved" else "",
                    "budget": "",
                }
            )

    sub_path = OUT_DIR / "PostAward - Subcontracts - DNNP.csv"
    change_path = OUT_DIR / "PostAward - Changes - DNNP.csv"
    write_csv(sub_path, SUB_COLUMNS, subcontracts)
    write_csv(change_path, CHANGE_COLUMNS, changes)

    raw_closeouts = {}
    for record in subcontracts:
        raw = record.get("_closeout_raw") or "(blank)"
        if not record["closeoutstatus"]:
            raw_closeouts[raw] = raw_closeouts.get(raw, 0) + 1
    print("subcontracts", len(subcontracts))
    print("changes", len(changes))
    print("sheets with change lines", sheets_with_changes)
    print("sheets added outside summary", unmatched_sheets)
    print("unmapped closeout", raw_closeouts)
    print("unmapped change status", unknown_status)
    print("unmapped change type", unknown_type)
    print("wrote", sub_path)
    print("wrote", change_path)


if __name__ == "__main__":
    main()
